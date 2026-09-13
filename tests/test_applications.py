from openpyxl import Workbook

from app.config import BASE_DIR


def test_create_and_get_application(client, sample_payload):
    created = client.post("/api/applications", json=sample_payload)
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["app_name"] == "AudioTo-txt"
    assert body["door"] == 5222
    assert body["docker"] == "Não"
    assert body["plan"] == "work"
    assert body["language"] == "Python"

    fetched = client.get(f"/api/applications/{body['id']}")
    assert fetched.status_code == 200
    assert fetched.json()["path"] == sample_payload["path"]


def test_list_filter_and_search(client, sample_payload):
    client.post("/api/applications", json=sample_payload)
    client.post(
        "/api/applications",
        json={
            **sample_payload,
            "app_name": "Portal",
            "path": "M:\\site",
            "door": 8080,
            "language": "HTML",
            "docker": "Sim",
            "plan": "personal",
            "nginx": "",
        },
    )
    work = client.get("/api/applications", params={"plan": "work"})
    assert work.status_code == 200
    assert work.json()["total"] == 1
    python = client.get("/api/applications", params={"language": "Python"})
    assert python.json()["total"] == 1
    docker_yes = client.get("/api/applications", params={"docker": "Sim"})
    assert docker_yes.json()["total"] == 1
    search = client.get("/api/applications", params={"q": "5222"})
    assert search.json()["total"] == 1


def test_update_and_delete_application(client, sample_payload):
    created = client.post("/api/applications", json=sample_payload).json()
    updated = client.put(
        f"/api/applications/{created['id']}",
        json={**sample_payload, "app_name": "AudioTo-txt-v2", "door": 5223},
    )
    assert updated.status_code == 200
    assert updated.json()["app_name"] == "AudioTo-txt-v2"
    assert updated.json()["door"] == 5223

    deleted = client.delete(f"/api/applications/{created['id']}")
    assert deleted.status_code == 200
    missing = client.get(f"/api/applications/{created['id']}")
    assert missing.status_code == 404


def test_bulk_delete_applications(client, sample_payload):
    first = client.post("/api/applications", json=sample_payload).json()
    second = client.post(
        "/api/applications",
        json={**sample_payload, "app_name": "Portal", "path": r"M:\site", "door": 8080},
    ).json()
    third = client.post(
        "/api/applications",
        json={**sample_payload, "app_name": "Keep", "path": r"M:\keep", "door": 9090},
    ).json()

    response = client.post(
        "/api/applications/bulk-delete",
        json={"ids": [first["id"], second["id"], 999999]},
    )
    assert response.status_code == 200, response.text
    assert response.json()["deleted"] == 2

    listed = client.get("/api/applications")
    assert listed.json()["total"] == 1
    assert listed.json()["items"][0]["id"] == third["id"]

    empty = client.post("/api/applications/bulk-delete", json={"ids": []})
    assert empty.status_code == 422

    missing = client.post("/api/applications/bulk-delete", json={"ids": [first["id"]]})
    assert missing.status_code == 404


def test_duplicate_door_is_rejected(client, sample_payload):
    assert client.post("/api/applications", json=sample_payload).status_code == 201
    conflict = client.post("/api/applications", json={**sample_payload, "app_name": "Outra"})
    assert conflict.status_code == 422
    assert "porta" in conflict.json()["detail"]["message"].lower()


def test_invalid_door_is_rejected(client, sample_payload):
    invalid = client.post("/api/applications", json={**sample_payload, "door": 70000})
    assert invalid.status_code == 422


def test_dashboard_counts(client, sample_payload):
    client.post("/api/applications", json=sample_payload)
    stats = client.get("/api/dashboard")
    assert stats.status_code == 200
    body = stats.json()
    assert body["total"] == 1
    assert body["work"] == 1
    assert body["python"] == 1
    assert body["cpp"] == 0
    assert body["docker_yes"] == 0
    assert body["nginx_yes"] == 1


def test_pages_render(client):
    assert client.get("/").status_code == 200
    assert client.get("/applications").status_code == 200
    assert client.get("/settings").status_code == 200
    assert client.get("/health").status_code == 200


def test_pages_render_under_doors_prefix(client):
    assert client.get("/doors/").status_code == 200
    assert client.get("/doors/applications").status_code == 200
    assert client.get("/doors/settings").status_code == 200
    assert client.get("/doors/health").status_code == 200
    assert client.get("/doors/api/dashboard").status_code == 200
    assert client.get("/Doors/").status_code == 200
    assert client.get("/Doors/api/dashboard").status_code == 200


def test_remote_prefix_is_injected_in_html(client):
    local = client.get("/").text
    assert 'href="/static/css/app.css' in local
    assert 'href="/doors/static/css/app.css' not in local

    remote = client.get("/doors/").text
    assert 'href="/doors/static/css/app.css' in remote
    assert 'data-app-root="/doors"' in remote
    assert 'data-stat="cpp"' in remote

    via_header = client.get("/", headers={"X-Forwarded-Prefix": "/doors"}).text
    assert 'href="/doors/static/css/app.css' in via_header

    via_upper = client.get("/", headers={"X-Forwarded-Prefix": "/Doors"}).text
    assert 'href="/doors/static/css/app.css' in via_upper
    assert 'data-app-root="/doors"' in via_upper


def test_create_language(client):
    created = client.post("/api/languages", json={"name": "JavaScript"})
    assert created.status_code == 201
    names = [item["name"] for item in client.get("/api/languages").json()]
    assert "JavaScript" in names
    assert "Python" in names
    assert "HTML" in names
    assert "C++" in names


def test_import_original_excel(client):
    xlsx = BASE_DIR / "Doors.xlsx"
    if not xlsx.exists():
        return
    with xlsx.open("rb") as handle:
        response = client.post("/api/import/xlsx", files={"file": ("Doors.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200
    assert response.json()["created"] >= 1
    listed = client.get("/api/applications", params={"q": "AudioTo-txt"})
    assert listed.json()["total"] >= 1


def _write_seed_xlsx(path, rows_by_sheet):
    workbook = Workbook()
    first = True
    for sheet_name, rows in rows_by_sheet.items():
        if first:
            sheet = workbook.active
            sheet.title = sheet_name
            first = False
        else:
            sheet = workbook.create_sheet(sheet_name)
        sheet["A1"] = "APP NAME"
        sheet["B1"] = "PATH"
        sheet["C1"] = "DOOR"
        sheet["D1"] = "LANGUAGE"
        sheet["E1"] = "NGINX"
        sheet["F1"] = "DOCKER"
        for index, row in enumerate(rows, start=3):
            for col, value in enumerate(row, start=1):
                sheet.cell(index, col, value)
    workbook.save(path)


def test_import_skips_empty_door_and_keeps_valid_row(client, tmp_path):
    seed = tmp_path / "seed.xlsx"
    _write_seed_xlsx(
        seed,
        {
            "Loterries": [
                ["SemPorta", r"D:\Loterias\x", None, "HTML", "", "Não"],
                ["LotteryLab", r"D:\Loterias\LotteryLab", 8082, "Python", "", "Não"],
            ]
        },
    )
    with seed.open("rb") as handle:
        response = client.post(
            "/api/import/xlsx",
            files={"file": ("seed.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert response.status_code == 200
    body = response.json()
    assert body["created"] == 1
    listed = client.get("/api/applications", params={"plan": "lotteries"})
    assert listed.json()["total"] == 1
    assert listed.json()["items"][0]["app_name"] == "LotteryLab"
    assert listed.json()["items"][0]["door"] == 8082


def test_import_is_idempotent_and_lotteries_win_door_conflict(client, tmp_path):
    seed = tmp_path / "seed.xlsx"
    _write_seed_xlsx(
        seed,
        {
            "Work": [["gifConverter", r"I:\apps\gifConverter", 5555, "Python", "", "Não"]],
            "Loterries": [
                ["LoteriasExtras-Conferencias-Lotofacil", r"D:\Loterias\LoteriasExtras\conferencias\lotofacil", 5555, "Python", "", "Não"]
            ],
        },
    )
    files = {"file": ("seed.xlsx", seed.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    first = client.post("/api/import/xlsx", files=files)
    second = client.post("/api/import/xlsx", files={"file": ("seed.xlsx", seed.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert first.status_code == 200
    assert first.json()["created"] == 1
    assert second.status_code == 200
    assert second.json()["created"] == 0
    listed = client.get("/api/applications", params={"q": "5555"})
    assert listed.json()["total"] == 1
    assert listed.json()["items"][0]["plan"] == "lotteries"
    assert listed.json()["items"][0]["app_name"] == "LoteriasExtras-Conferencias-Lotofacil"


def test_import_populada_persists_lotteries(client):
    xlsx = BASE_DIR / "Doors-populada.xlsx"
    if not xlsx.exists():
        return
    with xlsx.open("rb") as handle:
        response = client.post(
            "/api/import/xlsx",
            files={"file": ("Doors-populada.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert response.status_code == 200
    body = response.json()
    assert body["created"] >= 38
    lotteries = client.get("/api/applications", params={"plan": "lotteries", "page_size": 100})
    assert lotteries.json()["total"] == 38
    names = {item["app_name"] for item in lotteries.json()["items"]}
    assert "LotteryLab" in names
    assert "LoteriasPosicao-Central" in names
    assert client.get("/api/applications", params={"q": "8082"}).json()["total"] == 1
    again = client.post(
        "/api/import/xlsx",
        files={"file": ("Doors-populada.xlsx", xlsx.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert again.json()["created"] == 0
