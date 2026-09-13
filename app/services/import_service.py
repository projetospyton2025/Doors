from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.schemas.application import ApplicationCreate
from app.services.application_service import ApplicationService
from app.utils.constants import EXCEL_PLAN_ALIASES
from app.utils.validators import ValidationError

PLACEHOLDERS = {"", "selecione", None}
LOTTERY_SHEETS = {"loterries", "lotteries"}


class ImportService:
    def __init__(self, db: Session):
        self.service = ApplicationService(db)

    def import_xlsx(self, path: Path) -> dict:
        workbook = load_workbook(path, data_only=True)
        created = 0
        skipped = 0
        errors: list[str] = []

        for sheet_name in self._sheet_order(workbook.sheetnames):
            if sheet_name.strip().upper() == "DADOS":
                continue
            plan_key = EXCEL_PLAN_ALIASES.get(sheet_name.strip().lower())
            if not plan_key:
                skipped += 1
                continue
            sheet = workbook[sheet_name]
            for row in sheet.iter_rows(min_row=3, values_only=True):
                if not row or all(self._empty(value) for value in row):
                    continue
                app_name, path_value, door, language, nginx, docker, github, drive, *_ = (
                    list(row) + [None] * 8
                )[:8]
                if self._empty(app_name) and self._empty(path_value) and self._empty(door):
                    skipped += 1
                    continue
                if self._empty(door):
                    skipped += 1
                    continue
                if self._door_taken(door):
                    skipped += 1
                    continue
                try:
                    self.service.create(
                        ApplicationCreate(
                            app_name=str(app_name or ""),
                            path=str(path_value or ""),
                            door=door,
                            language=str(language or ""),
                            nginx=None if self._empty(nginx) else str(nginx),
                            docker=str(docker or ""),
                            plan=plan_key,
                            github=None if self._empty(github) else str(github),
                            drive=None if self._empty(drive) else str(drive),
                        )
                    )
                    created += 1
                except (ValidationError, LookupError) as exc:
                    errors.append(f"{sheet_name}: {exc}")
                    skipped += 1
        return {"created": created, "skipped": skipped, "errors": errors}

    def _sheet_order(self, sheetnames: list[str]) -> list[str]:
        lottery = [name for name in sheetnames if name.strip().lower() in LOTTERY_SHEETS]
        others = [name for name in sheetnames if name.strip().lower() not in LOTTERY_SHEETS]
        return lottery + others

    def _door_taken(self, door: object) -> bool:
        try:
            door_int = int(str(door).strip())
        except (TypeError, ValueError):
            return False
        return self.service.applications.get_by_door(door_int) is not None

    def _empty(self, value: object) -> bool:
        if value is None:
            return True
        return str(value).strip().lower() in PLACEHOLDERS
