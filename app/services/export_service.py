import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.application import ApplicationOut
from app.utils.constants import EXPORT_COLUMNS, PLANS


class ExportService:
    def rows(self, items: list[ApplicationOut]) -> list[list[object]]:
        return [
            [
                item.plan_label,
                item.app_name,
                item.path,
                item.door,
                item.language,
                item.nginx or "",
                item.docker,
                item.github or "",
                item.drive or "",
            ]
            for item in items
        ]

    def csv_bytes(self, items: list[ApplicationOut]) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer, lineterminator="\n")
        writer.writerow(EXPORT_COLUMNS)
        writer.writerows(self.rows(items))
        return buffer.getvalue().encode("utf-8-sig")

    def xlsx_bytes(self, items: list[ApplicationOut], *, split_by_plan: bool = True) -> bytes:
        workbook = Workbook()
        header_fill = PatternFill("solid", fgColor="2A303B")
        header_font = Font(color="FFFFFF", bold=True)
        if split_by_plan and items:
            first = True
            for plan in PLANS:
                plan_items = [item for item in items if item.plan == plan["key"]]
                if first:
                    sheet = workbook.active
                    sheet.title = plan["label"]
                    first = False
                else:
                    sheet = workbook.create_sheet(plan["label"])
                self._write_sheet(sheet, plan_items, header_fill, header_font)
            leftover = [item for item in items if item.plan not in {p["key"] for p in PLANS}]
            if leftover:
                sheet = workbook.create_sheet("Outros")
                self._write_sheet(sheet, leftover, header_fill, header_font)
        else:
            sheet = workbook.active
            sheet.title = "Doors"
            self._write_sheet(sheet, items, header_fill, header_font)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def filename(self, extension: str) -> str:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        return f"doors-{stamp}.{extension}"

    def _write_sheet(self, sheet, items: list[ApplicationOut], header_fill, header_font) -> None:
        sheet.merge_cells("G1:H1")
        sheet["A1"] = "APP NAME"
        sheet["B1"] = "PATH"
        sheet["C1"] = "DOOR"
        sheet["D1"] = "LANGUAGE"
        sheet["E1"] = "NGINX"
        sheet["F1"] = "DOCKER"
        sheet["G1"] = "ACCOUNT"
        sheet["G2"] = "GITHUB"
        sheet["H2"] = "DRIVE"
        for cell in ("A1", "B1", "C1", "D1", "E1", "F1", "G1", "G2", "H2"):
            sheet[cell].fill = header_fill
            sheet[cell].font = header_font
            sheet[cell].alignment = Alignment(horizontal="center", vertical="center")
        for index, item in enumerate(items, start=3):
            sheet.cell(index, 1, item.app_name)
            sheet.cell(index, 2, item.path)
            sheet.cell(index, 3, item.door)
            sheet.cell(index, 4, item.language)
            sheet.cell(index, 5, item.nginx or "")
            sheet.cell(index, 6, item.docker)
            sheet.cell(index, 7, item.github or "")
            sheet.cell(index, 8, item.drive or "")
        widths = [26, 46, 12, 14, 42, 12, 32, 32]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
