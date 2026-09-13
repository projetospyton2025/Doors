"""Sincroniza o links.xlsx (legenda por cor + coluna Remoto) com a plataforma."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.schemas.application import ApplicationCreate
from app.services.application_service import ApplicationService
from app.repositories.application_repository import ApplicationRepository
from app.utils.constants import (
    LANGUAGE_HTML,
    LANGUAGE_PYTHON,
    is_satellite_app,
)
from app.utils.validators import ValidationError

# Legenda do links.xlsx (cores da fonte)
COLOR_TO_PLAN = {
    "00B0F0": "personal",  # AZUL = Particular
    "7030A0": "lotteries",  # ROXO = Loterias
    "FFC000": "work",  # LARANJA = Trabalho/Work
}

SKIP_NAMES = {
    "resultado",
    "legenda",
    "roxo",
    "azul",
    "laranja",
}


def _font_rgb(cell) -> str | None:
    color = cell.font.color if cell.font else None
    if color is None:
        return None
    if color.type == "rgb" and color.rgb:
        return str(color.rgb).upper()[-6:]
    return None


def _parse_door(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    if not text:
        return None
    # http://localhost:8086/5575 → usa o número final (porta da app)
    tail = re.search(r":\d+/(\d{3,5})\s*$", text)
    if tail:
        return int(tail.group(1))
    ports = re.findall(r":(\d{2,5})", text)
    if ports:
        return int(ports[-1])
    if text.isdigit():
        return int(text)
    return None


def _guess_language(app_name: str, path: str, ambiente: str | None) -> str:
    blob = f"{app_name} {path} {ambiente or ''}".lower()
    if any(token in blob for token in ("html", "javascript", "css", "netlify", "portifolio", "portfolio")):
        return LANGUAGE_HTML
    return LANGUAGE_PYTHON


def parse_links_xlsx(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows: list[dict] = []
    for index in range(2, (sheet.max_row or 1) + 1):
        name_cell = sheet.cell(index, 1)
        app_name = name_cell.value
        if app_name is None:
            continue
        app_name = str(app_name).strip()
        if not app_name or app_name.lower() in SKIP_NAMES:
            continue
        if is_satellite_app(app_name):
            continue

        door = _parse_door(sheet.cell(index, 2).value)
        if door is None:
            continue

        ambiente = sheet.cell(index, 4).value
        path_value = sheet.cell(index, 6).value
        remoto = sheet.cell(index, 7).value
        path_text = str(path_value).strip() if path_value else ""
        nginx = str(remoto).strip() if remoto else None
        if nginx == "":
            nginx = None

        rgb = _font_rgb(name_cell)
        plan = COLOR_TO_PLAN.get(rgb or "")
        # Audio-To-Txt (fonte tema) fica com Particular, junto dos azuis
        if plan is None:
            plan = "personal"

        rows.append(
            {
                "app_name": app_name,
                "door": door,
                "path": path_text or f"(links.xlsx) {app_name}",
                "nginx": nginx,
                "plan": plan,
                "language": _guess_language(app_name, path_text, str(ambiente) if ambiente else None),
                "docker": "Não",
            }
        )
    return rows


def is_links_workbook(path: Path) -> bool:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        headers = [str(sheet.cell(1, col).value or "").strip().lower() for col in range(1, 8)]
        return headers[:3] == ["resultado", "porta", "acesso"] and "remoto" in headers
    finally:
        workbook.close()


class LinksSyncService:
    def __init__(self, db: Session):
        self.db = db
        self.service = ApplicationService(db)
        self.repo = ApplicationRepository(db)

    def sync_file(self, path: Path) -> dict:
        items = parse_links_xlsx(path)
        created = 0
        updated = 0
        skipped = 0
        errors: list[str] = []

        for item in items:
            existing = self.repo.get_by_door(item["door"])
            if existing is None:
                try:
                    self.service.create(
                        ApplicationCreate(
                            app_name=item["app_name"],
                            path=item["path"],
                            door=item["door"],
                            language=item["language"],
                            nginx=item["nginx"],
                            docker=item["docker"],
                            plan=item["plan"],
                            github=None,
                            drive=None,
                        )
                    )
                    created += 1
                except (ValidationError, LookupError) as exc:
                    errors.append(f"{item['app_name']}: {exc}")
                    skipped += 1
                continue

            changed = False
            # Preenche Remoto se vazio; se já houver, atualiza quando o Excel trouxer URL
            if item["nginx"] and existing.nginx != item["nginx"]:
                existing.nginx = item["nginx"]
                changed = True
            if item["path"] and (not existing.path or existing.path.startswith("(links.xlsx)")):
                existing.path = item["path"]
                changed = True
            # Não altera plano/nome de registros já existentes sem necessidade
            if changed:
                self.db.commit()
                updated += 1
            else:
                skipped += 1

        return {
            "created": created,
            "updated": updated,
            "skipped": skipped,
            "errors": errors,
            "total_in_file": len(items),
        }
